from __future__ import annotations

import hashlib
from collections import Counter
from collections.abc import Iterable
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def _safe_float(value: Any, default: float = 0.0) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def _iter_shapes(shapes: Iterable[Any]) -> Iterable[Any]:
    for shape in shapes:
        yield shape
        if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shape.shapes)


def _bbox(shape: Any) -> tuple[int, int, int, int]:
    left = int(getattr(shape, "left", 0))
    top = int(getattr(shape, "top", 0))
    width = max(0, int(getattr(shape, "width", 0)))
    height = max(0, int(getattr(shape, "height", 0)))
    return left, top, left + width, top + height


def _area(box: tuple[int, int, int, int]) -> int:
    return max(0, box[2] - box[0]) * max(0, box[3] - box[1])


def _intersection_ratio(a: tuple[int, int, int, int], b: tuple[int, int, int, int]) -> float:
    x1, y1 = max(a[0], b[0]), max(a[1], b[1])
    x2, y2 = min(a[2], b[2]), min(a[3], b[3])
    intersection = max(0, x2 - x1) * max(0, y2 - y1)
    denominator = min(_area(a), _area(b))
    return intersection / denominator if denominator else 0.0


def _shape_kind(shape: Any) -> str:
    shape_type = getattr(shape, "shape_type", None)
    mapping = {
        MSO_SHAPE_TYPE.PICTURE: "picture",
        MSO_SHAPE_TYPE.CHART: "chart",
        MSO_SHAPE_TYPE.TABLE: "table",
        MSO_SHAPE_TYPE.GROUP: "group",
        MSO_SHAPE_TYPE.MEDIA: "media",
        MSO_SHAPE_TYPE.TEXT_BOX: "text_box",
        MSO_SHAPE_TYPE.PLACEHOLDER: "placeholder",
        MSO_SHAPE_TYPE.AUTO_SHAPE: "auto_shape",
    }
    return mapping.get(shape_type, "other")


def _text_metrics(shape: Any) -> tuple[int, list[float], list[str]]:
    if not getattr(shape, "has_text_frame", False):
        return 0, [], []
    text = shape.text or ""
    sizes: list[float] = []
    fonts: list[str] = []
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if run.font.size is not None:
                sizes.append(round(run.font.size.pt, 2))
            if run.font.name:
                fonts.append(run.font.name)
    return len(text.strip()), sizes, fonts


def inspect_pptx(path: str | Path, *, overlap_threshold: float = 0.25) -> dict[str, Any]:
    source = Path(path)
    if not source.is_file():
        raise FileNotFoundError(source)
    if source.suffix.lower() != ".pptx":
        raise ValueError("The native inspector currently supports .pptx files only")

    prs = Presentation(str(source))
    width, height = int(prs.slide_width), int(prs.slide_height)
    type_counts: Counter[str] = Counter()
    font_counts: Counter[str] = Counter()
    all_font_sizes: list[float] = []
    slide_reports: list[dict[str, Any]] = []
    total_chars = 0
    total_notes = 0
    total_overflow = 0
    total_overlap = 0
    low_dpi_images = 0

    for slide_number, slide in enumerate(prs.slides, start=1):
        shapes = list(_iter_shapes(slide.shapes))
        visible_boxes: list[tuple[Any, tuple[int, int, int, int]]] = []
        overflow_risks: list[dict[str, Any]] = []
        overlap_risks: list[dict[str, Any]] = []
        image_dpi: list[dict[str, Any]] = []
        char_count = 0

        for shape in shapes:
            kind = _shape_kind(shape)
            type_counts[kind] += 1
            chars, sizes, fonts = _text_metrics(shape)
            char_count += chars
            all_font_sizes.extend(sizes)
            font_counts.update(fonts)

            box = _bbox(shape)
            if box[0] < 0 or box[1] < 0 or box[2] > width or box[3] > height:
                overflow_risks.append({"shape": shape.name, "kind": kind, "bbox_emu": box})

            # Exclude near-full-slide shapes because backgrounds commonly overlap everything.
            if _area(box) and _area(box) < width * height * 0.90 and kind not in {"group", "other"}:
                visible_boxes.append((shape, box))

            if kind == "picture":
                image = shape.image
                px_w, px_h = image.size
                inch_w = max(_safe_float(shape.width) / 914400, 1e-9)
                inch_h = max(_safe_float(shape.height) / 914400, 1e-9)
                effective_dpi = round(min(px_w / inch_w, px_h / inch_h), 1)
                image_dpi.append({"shape": shape.name, "effective_dpi": effective_dpi})
                if effective_dpi < 96:
                    low_dpi_images += 1

        for index, (left_shape, left_box) in enumerate(visible_boxes):
            for right_shape, right_box in visible_boxes[index + 1 :]:
                ratio = _intersection_ratio(left_box, right_box)
                if ratio >= overlap_threshold:
                    overlap_risks.append(
                        {
                            "shape_a": left_shape.name,
                            "shape_b": right_shape.name,
                            "intersection_over_smaller": round(ratio, 4),
                        }
                    )

        notes_text = ""
        if slide.has_notes_slide:
            notes_text = slide.notes_slide.notes_text_frame.text.strip()
        if notes_text:
            total_notes += 1

        total_chars += char_count
        total_overflow += len(overflow_risks)
        total_overlap += len(overlap_risks)
        slide_reports.append(
            {
                "slide_number": slide_number,
                "shape_count": len(shapes),
                "character_count": char_count,
                "overflow_risks": overflow_risks,
                "overlap_risks": overlap_risks,
                "images": image_dpi,
                "has_speaker_notes": bool(notes_text),
            }
        )

    sha256 = hashlib.sha256(source.read_bytes()).hexdigest()
    slide_count = len(prs.slides)
    editable_kinds = {"text_box", "placeholder", "auto_shape", "chart", "table"}
    editable_objects = sum(type_counts[kind] for kind in editable_kinds)
    content_objects = editable_objects + type_counts["picture"] + type_counts["media"]

    return {
        "schema_version": "0.1.0",
        "artifact": {
            "filename": source.name,
            "sha256": sha256,
            "size_bytes": source.stat().st_size,
            "format": "pptx",
        },
        "presentation": {
            "slide_count": slide_count,
            "canvas_width_emu": width,
            "canvas_height_emu": height,
            "aspect_ratio": round(width / height, 4) if height else None,
        },
        "summary": {
            "shape_counts": dict(sorted(type_counts.items())),
            "editable_object_ratio": round(editable_objects / content_objects, 4)
            if content_objects
            else None,
            "total_characters": total_chars,
            "mean_characters_per_slide": round(total_chars / slide_count, 2) if slide_count else 0,
            "minimum_explicit_font_size_pt": min(all_font_sizes) if all_font_sizes else None,
            "median_explicit_font_size_pt": _median(all_font_sizes),
            "font_families": dict(font_counts.most_common()),
            "slides_with_speaker_notes": total_notes,
            "overflow_risk_count": total_overflow,
            "overlap_risk_count": total_overlap,
            "low_effective_dpi_image_count": low_dpi_images,
        },
        "slides": slide_reports,
        "limitations": [
            "Overlap and overflow are geometric risk signals and require visual confirmation.",
            "Font sizes are reported only where explicitly stored on text runs.",
            "Rendering-dependent text clipping is not detected in v0.1.",
        ],
    }


def _median(values: list[float]) -> float | None:
    if not values:
        return None
    ordered = sorted(values)
    middle = len(ordered) // 2
    if len(ordered) % 2:
        return ordered[middle]
    return round((ordered[middle - 1] + ordered[middle]) / 2, 2)
