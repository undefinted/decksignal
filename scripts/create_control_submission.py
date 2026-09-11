"""Create a clearly labelled non-AI control deck for pipeline smoke tests."""

from __future__ import annotations

import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "results" / "raw" / "pilot-control" / "control" / "opb-business-001"


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    deck = OUT / "deck.pptx"
    prs = Presentation()
    for title, body in [
        ("DeckSignal control artifact", "Deterministic smoke-test deck; not an AI product result."),
        ("Pilot workflow", "Manifest → native PPTX inspection → benchmark adapters → reviewed evidence."),
    ]:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = title
        box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(2))
        paragraph = box.text_frame.paragraphs[0]
        paragraph.text = body
        paragraph.font.size = Pt(24)
    prs.save(deck)
    digest = hashlib.sha256(deck.read_bytes()).hexdigest()
    manifest = {
        "schema_version": "0.1.0",
        "submission_id": "control-opb-business-001",
        "task_id": "opb-business-001",
        "track": "control",
        "product": {"name": "DeckSignal control", "provider": "DeckSignal", "version": "local", "plan": "none", "url": None},
        "run": {"started_at": datetime.now(timezone.utc).isoformat(), "completed_at": datetime.now(timezone.utc).isoformat(), "attempt": 1, "locale": "zh-CN", "manual_edits": False, "notes": "Control only; excluded from rankings."},
        "artifact": {"format": "pptx", "path": "deck.pptx", "sha256": digest, "redistribution": "allowed"},
        "scores": {},
        "metrics_path": None,
    }
    (OUT / "submission.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(OUT.relative_to(ROOT))


if __name__ == "__main__":
    main()
