from pptx import Presentation
from pptx.util import Inches, Pt

from openpptbench.inspect import inspect_pptx


def test_inspector_reports_structure_and_risks(tmp_path):
    path = tmp_path / "sample.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    first = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    run = first.text_frame.paragraphs[0].add_run()
    run.text = "OpenPPTBench"
    run.font.size = Pt(24)
    second = slide.shapes.add_textbox(Inches(2), Inches(1), Inches(3), Inches(1))
    second.text = "overlap"
    slide.shapes.add_textbox(Inches(-0.2), Inches(2), Inches(1), Inches(1)).text = "outside"
    prs.save(path)

    result = inspect_pptx(path)

    assert result["presentation"]["slide_count"] == 1
    assert result["summary"]["minimum_explicit_font_size_pt"] == 24
    assert result["summary"]["overflow_risk_count"] >= 1
    assert result["summary"]["overlap_risk_count"] >= 1
    assert result["artifact"]["sha256"]

